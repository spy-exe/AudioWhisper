import re
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

_BOLD = re.compile(r"\*\*(.+?)\*\*")


def inline_runs(text: str) -> list[tuple[str, bool]]:
    """Split the only inline markup our Markdown uses (**bold**) into (text, bold) runs."""
    runs, cursor = [], 0
    for match in _BOLD.finditer(text):
        if match.start() > cursor:
            runs.append((text[cursor:match.start()], False))
        runs.append((match.group(1), True))
        cursor = match.end()
    if cursor < len(text):
        runs.append((text[cursor:], False))
    return [run for run in runs if run[0]]


def document_blocks(markdown: str) -> list[tuple[str, str]]:
    """Markdown lines as (kind, text): title, heading, bullet, field or paragraph. Tables become
    "Campo: valor" fields (header and separator rows dropped) instead of literal pipes."""
    blocks, in_table = [], False
    for line in markdown.splitlines():
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            separator = all(cell and set(cell) <= set("-: ") for cell in cells)
            if in_table and not separator and len(cells) >= 2:
                blocks.append(("field", f"**{cells[0]}:** {' · '.join(cells[1:])}"))
            in_table = True
            continue
        in_table = False
        if line.startswith("# "):
            blocks.append(("title", line[2:]))
        elif line.startswith("## "):
            blocks.append(("heading", line[3:]))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:]))
        elif stripped:
            blocks.append(("paragraph", stripped))
        else:
            blocks.append(("blank", ""))
    return blocks


def export_document(fmt: str, title: str, markdown: str) -> bytes:
    """Render plain, escaped document content without evaluating HTML or fetching assets."""
    output = BytesIO()
    blocks = document_blocks(markdown)
    if fmt == "docx":
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(.8)
        normal = document.styles["Normal"]
        normal.font.name, normal.font.size = "Calibri", Pt(11)
        normal.font.color.rgb = RGBColor.from_string("253442")
        normal.paragraph_format.space_after = Pt(8)
        document.core_properties.title = title
        for kind, text in blocks:
            if kind == "title":
                document.add_heading(text, 0)
            elif kind == "heading":
                document.add_heading(text, 1)
            elif kind in ("bullet", "field", "paragraph"):
                paragraph = document.add_paragraph(style="List Bullet" if kind == "bullet" else None)
                for run_text, bold in inline_runs(text):
                    paragraph.add_run(run_text).bold = bold
        section.footer.paragraphs[0].text = "AudioWhisper Studio · Processamento local"
        document.save(output)
    elif fmt == "pdf":
        import reportlab
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        fonts = Path(reportlab.__file__).parent / "fonts"
        if "AudioWhisper" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("AudioWhisper", str(fonts / "Vera.ttf")))
            pdfmetrics.registerFont(TTFont("AudioWhisper-Bold", str(fonts / "VeraBd.ttf")))
            pdfmetrics.registerFontFamily("AudioWhisper", normal="AudioWhisper", bold="AudioWhisper-Bold",
                                          italic="AudioWhisper", boldItalic="AudioWhisper-Bold")
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = "AudioWhisper"
        styles["Normal"].leading = 16
        styles["Normal"].spaceAfter = 8
        styles["Title"].textColor = colors.HexColor("#163948")
        bullet = styles["Bullet"]
        bullet.leftIndent, bullet.bulletIndent, bullet.leading = 14, 4, 16
        story = []
        for kind, text in blocks:
            if kind == "blank":
                story.append(Spacer(1, 6))
                continue
            markup = "".join(f"<b>{escape(t)}</b>" if bold else escape(t) for t, bold in inline_runs(text))
            style = {"title": "Title", "heading": "Heading2", "bullet": "Bullet"}.get(kind, "Normal")
            story.append(Paragraph(markup, styles[style], bulletText="•" if kind == "bullet" else None))
        def footer(canvas, document):
            canvas.setFont("AudioWhisper", 8)
            canvas.setFillColor(colors.HexColor("#526879"))
            canvas.drawString(48, 30, "AudioWhisper Studio · Processamento local")
            canvas.drawRightString(A4[0] - 48, 30, str(document.page))
        SimpleDocTemplate(output, pagesize=A4, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=54, title=title).build(story, onFirstPage=footer, onLaterPages=footer)
    else:
        raise ValueError("Formato de documento não suportado")
    return output.getvalue()
